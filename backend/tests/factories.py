import io

from docx import Document as DocxDocument
from fpdf import FPDF

SAMPLE_RESUME = [
    "Jane Doe - AI Engineer",
    "SUMMARY",
    "Backend engineer with 4 years of experience building Python services and LLM applications.",
    "EXPERIENCE",
    "Acme Corp - Software Engineer (2021-2025)",
    "- Built a FastAPI service handling 2M requests per day with PostgreSQL and Redis.",
    "- Designed a RAG pipeline using pgvector and hybrid search for support documents.",
    "SKILLS",
    "Python, FastAPI, PostgreSQL, Docker, Kubernetes, scikit-learn, LangChain",
]

SAMPLE_JD = (
    "We are hiring a Machine Learning Engineer to build retrieval-augmented generation "
    "systems. You will design Python services with FastAPI, work with PostgreSQL and "
    "vector search, and own evaluation of LLM output quality in production."
)


def make_pdf(lines=None, pages: int = 1) -> bytes:
    lines = SAMPLE_RESUME if lines is None else lines
    pdf = FPDF()
    for _ in range(pages):
        pdf.add_page()
        pdf.set_font("Helvetica", size=11)
        for line in lines:
            pdf.multi_cell(0, 6, line, new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


def make_blank_pdf() -> bytes:
    pdf = FPDF()
    pdf.add_page()  # no text at all: behaves like an image-only scan
    return bytes(pdf.output())


def make_docx(paragraphs=None, bullets=(), table_rows=()) -> bytes:
    d = DocxDocument()
    for p in SAMPLE_RESUME if paragraphs is None else paragraphs:
        d.add_paragraph(p)
    for b in bullets:
        d.add_paragraph(b, style="List Bullet")
    if table_rows:
        table = d.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for r, row in enumerate(table_rows):
            for c, value in enumerate(row):
                table.cell(r, c).text = value
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()