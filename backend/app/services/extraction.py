import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass
from pathlib import Path

from docx import Document as DocxDocument
from docx.table import Table
from docx.text.paragraph import Paragraph
import pymupdf
MAX_DOCX_UNCOMPRESSED = 50 * 1024 * 1024  # zip-bomb guard

MIME_BY_EXT = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


class DocumentError(Exception):
    status_code = 422

    def __init__(self, message: str):
        super().__init__(message)
        self.message = message


class UnsupportedFileType(DocumentError):
    status_code = 415


class FileTooLarge(DocumentError):
    status_code = 413


class UnreadableDocument(DocumentError):
    status_code = 422


@dataclass
class ExtractedDoc:
    text: str
    mime: str
    page_count: int | None


_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


def normalize_text(text: str) -> str:
    """Clean extraction noise without touching technical terms (C++, Node.js, k8s)."""
    text = unicodedata.normalize("NFKC", text)  # also expands ligatures like "fi"
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = _CONTROL.sub("", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _check_signature(ext: str, data: bytes) -> None:
    if ext == ".pdf":
        ok = b"%PDF-" in data[:1024]
        label = "PDF"
    else:
        ok = data[:4] == b"PK\x03\x04" and zipfile.is_zipfile(io.BytesIO(data))
        label = "Word (.docx)"
    if not ok:
        raise UnreadableDocument(f"This file doesn't look like a valid {label} document.")


def _pdf_text(data: bytes, max_pages: int) -> tuple[str, int]:
    try:
        with pymupdf.open(stream=data, filetype="pdf") as doc:
            if doc.needs_pass:
                raise UnreadableDocument(
                    "This PDF is password-protected. Remove the password and upload it again."
                )
            pages = doc.page_count
            if pages > max_pages:
                raise UnreadableDocument(
                    f"This PDF has {pages} pages; the limit is {max_pages} for a resume."
                )
            text = "\n".join(page.get_text("text") for page in doc)
            return text, pages
    except DocumentError:
        raise
    except Exception:
        raise UnreadableDocument("Could not read this PDF. It may be corrupted.") from None


def _docx_text(data: bytes) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if sum(i.file_size for i in z.infolist()) > MAX_DOCX_UNCOMPRESSED:
                raise UnreadableDocument("This document expands to an unsafe size.")
            if "word/document.xml" not in z.namelist():
                raise UnreadableDocument("This isn't a valid Word document.")
        doc = DocxDocument(io.BytesIO(data))
        lines: list[str] = []
        for child in doc.element.body.iterchildren():
            tag = child.tag.rsplit("}", 1)[-1]
            if tag == "p":
                para = Paragraph(child, doc)
                t = para.text.strip()
                if not t:
                    lines.append("")
                    continue
                style = (para.style.name if para.style is not None else "") or ""
                lines.append(f"• {t}" if style.startswith("List") else t)
            elif tag == "tbl":
                for row in Table(child, doc).rows:
                    cells: list[str] = []
                    for cell in row.cells:
                        ct = cell.text.strip().replace("\n", " ")
                        if ct and (not cells or cells[-1] != ct):  # merged cells repeat
                            cells.append(ct)
                    if cells:
                        lines.append(" | ".join(cells))
        return "\n".join(lines)
    except DocumentError:
        raise
    except Exception:
        raise UnreadableDocument("Could not read this Word document.") from None


def extract_text(
    filename: str, data: bytes, *, max_pages: int = 15, min_chars: int = 200
) -> ExtractedDoc:
    ext = Path(filename).suffix.lower()
    if ext not in MIME_BY_EXT:
        raise UnsupportedFileType("Only PDF and DOCX files are supported.")
    _check_signature(ext, data)

    page_count: int | None = None
    if ext == ".pdf":
        raw, page_count = _pdf_text(data, max_pages)
    else:
        raw = _docx_text(data)

    text = normalize_text(raw)
    if len(text) < min_chars:
        raise UnreadableDocument(
            f"We could only read {len(text)} characters from this file. It may be a scanned "
            "or image-based document (OCR isn't supported yet). Upload a text-based PDF "
            "or a DOCX."
        )
    return ExtractedDoc(text=text, mime=MIME_BY_EXT[ext], page_count=page_count)